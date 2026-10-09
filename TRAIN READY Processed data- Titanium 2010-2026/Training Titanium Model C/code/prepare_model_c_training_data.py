"""Prepare one monthly research dataset for the agreed Model C split."""
from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / 'Data'
OUT_DIR = ROOT / 'Train'
FREIGHT_PATH = DATA_DIR / 'bls_deep_sea_freight_ppi_monthly_1988-06_to_2026-08.csv'
BRENT_PATH = DATA_DIR / 'world_bank_pink_sheet_monthly_1960-01_to_2026-09.xlsx'
OUTPUT = OUT_DIR / 'model_c_freight_training_2005_2026.csv'
HORIZONS = (3, 6, 12)
FEATURES = [
    'freight_log_change_1m', 'freight_log_change_2m', 'freight_log_change_3m',
    'freight_log_change_6m', 'freight_log_change_12m',
    'freight_volatility_6m', 'freight_volatility_12m', 'brent_log_change_1m',
]
START = pd.Timestamp('2005-01-01')
WARMUP_START = pd.Timestamp('2004-01-01')
SELECTION_FREEZE = pd.Timestamp('2020-01-01')
PREPARED_DATE = '2026-10-09'


def check_months(frame, label):
    if frame.date.isna().any() or frame.date.duplicated().any():
        raise ValueError(f'{label}: missing or duplicate dates')
    periods = pd.PeriodIndex(frame.date, freq='M')
    expected = pd.period_range(periods.min(), periods.max(), freq='M')
    if not periods.equals(expected):
        raise ValueError(f'{label}: missing or out-of-order months')


def read_freight():
    frame = pd.read_csv(FREIGHT_PATH, dtype={'observation_date': str})
    if list(frame.columns) != ['observation_date', 'PCU483111483111']:
        raise ValueError('Unexpected freight CSV columns')
    frame = frame.rename(columns={'observation_date': 'date', 'PCU483111483111': 'freight_index'})
    frame.date = pd.to_datetime(frame.date, errors='raise')
    frame.freight_index = pd.to_numeric(frame.freight_index, errors='raise')
    frame = frame.sort_values('date').reset_index(drop=True)
    check_months(frame, 'Freight')
    if frame.freight_index.isna().any() or (frame.freight_index <= 0).any():
        raise ValueError('Freight values must be positive observations')
    return frame


def read_brent():
    workbook = load_workbook(BRENT_PATH, read_only=True, data_only=True)
    try:
        sheet = workbook['Monthly Prices']
        header_rows = list(sheet.iter_rows(min_row=1, max_row=10, values_only=True))
        matches = [(r, c) for r, row in enumerate(header_rows, 1) for c, value in enumerate(row)
                   if str(value).strip().lower() == 'crude oil, brent']
        if len(matches) != 1:
            raise ValueError('Cannot uniquely identify World Bank Brent column')
        header_row, column = matches[0]
        records = []
        for row in sheet.iter_rows(min_row=header_row+1, values_only=True):
            period = str(row[0])
            if len(period) == 7 and period[:4].isdigit() and period[4] == 'M' and period[5:].isdigit():
                records.append({'date': pd.Timestamp(int(period[:4]), int(period[5:]), 1),
                                'brent_usd_per_barrel': row[column]})
        frame = pd.DataFrame(records).sort_values('date').reset_index(drop=True)
        frame.brent_usd_per_barrel = pd.to_numeric(frame.brent_usd_per_barrel, errors='raise')
        check_months(frame, 'Brent')
        if frame.brent_usd_per_barrel.isna().any() or (frame.brent_usd_per_barrel <= 0).any():
            raise ValueError('Brent values must be positive observations')
        return frame
    finally:
        workbook.close()


def period_group(dates):
    return np.select([
        dates <= pd.Timestamp('2014-12-01'),
        dates <= pd.Timestamp('2019-12-01'),
        dates <= pd.Timestamp('2021-12-01'),
    ], ['initial_training', 'validation', 'disruption_test'], default='post_disruption')


def prepare(freight, brent):
    # Calendar-key merge and gap checks make shift(h) mean h calendar months.
    work = freight.merge(brent, on='date', how='inner', validate='one_to_one')
    work = work.loc[work.date >= WARMUP_START].copy().reset_index(drop=True)
    check_months(work, 'Joined warm-up history')
    if work.date.min() != WARMUP_START:
        raise ValueError('The 2004 warm-up history is incomplete')
    for lag in (1, 2, 3, 6, 12):
        work[f'freight_log_change_{lag}m'] = np.log(work.freight_index / work.freight_index.shift(lag))
    work['brent_log_change_1m'] = np.log(work.brent_usd_per_barrel / work.brent_usd_per_barrel.shift(1))
    for window in (6, 12):
        work[f'freight_volatility_{window}m'] = work.freight_log_change_1m.rolling(window, min_periods=window).std(ddof=0)

    # These are explicit prototype timing assumptions, not invented release dates.
    # Both source observations are treated as usable at the end of the next month.
    work['freight_available_date_assumed'] = work.date + pd.offsets.MonthEnd(2)
    work['brent_available_date_assumed'] = work.date + pd.offsets.MonthEnd(2)
    work['feature_available_date_assumed'] = work[['freight_available_date_assumed', 'brent_available_date_assumed']].max(axis=1)
    work['forecast_date_assumed'] = work.feature_available_date_assumed
    work['freight_release_date_verified'] = pd.NaT
    work['brent_release_date_verified'] = pd.NaT
    work['historical_release_dates_verified'] = False
    work['historical_data_vintages_verified'] = False
    work['availability_policy'] = 'ASSUMPTION: reference-month end plus one full calendar month'
    work['source_value_basis'] = 'downloaded historical snapshot; original historical vintages not verified'

    for h in HORIZONS:
        # Target month is retained even when the future observed value is unavailable.
        work[f'target_month_h{h}'] = work.date + pd.DateOffset(months=h)
        work[f'target_freight_index_h{h}'] = work.freight_index.shift(-h)
        work[f'target_freight_log_change_h{h}'] = np.log(work[f'target_freight_index_h{h}'] / work.freight_index)
        work[f'target_freight_pct_change_h{h}'] = np.expm1(work[f'target_freight_log_change_h{h}']) * 100
        work[f'target_available_date_assumed_h{h}'] = work[f'target_month_h{h}'] + pd.offsets.MonthEnd(2)
        work[f'target_observed_h{h}'] = work[f'target_freight_index_h{h}'].notna()

    frame = work.loc[work.date >= START].copy().reset_index(drop=True)
    frame.insert(1, 'period', frame.date.dt.strftime('%Y-%m'))
    frame.insert(2, 'year', frame.date.dt.year)
    frame.insert(3, 'month', frame.date.dt.month)
    frame['origin_period_group'] = period_group(frame.date)
    frame['forecast_issue_period_group_assumed'] = period_group(frame.forecast_date_assumed.dt.to_period('M').dt.to_timestamp())
    frame['optional_early_backtest_origin'] = frame.date.between('2012-01-01', '2014-12-01')
    frame['feature_complete'] = frame[FEATURES].notna().all(axis=1)
    frame['h12_reporting_only'] = True
    first_validation_asof = frame.loc[frame.origin_period_group.eq('validation'), 'forecast_date_assumed'].min()
    frame['model_selection_freeze_date'] = SELECTION_FREEZE
    frame['initial_validation_asof_date_assumed'] = first_validation_asof
    for h in HORIZONS:
        complete = frame.feature_complete & frame[f'target_observed_h{h}']
        frame[f'research_training_eligible_h{h}'] = complete
        frame[f'initial_training_fit_eligible_h{h}'] = (
            complete & frame.origin_period_group.eq('initial_training')
            & (frame[f'target_available_date_assumed_h{h}'] <= first_validation_asof))
        frame[f'validation_selection_eligible_h{h}'] = (
            complete & frame.origin_period_group.eq('validation')
            & (frame.forecast_date_assumed < SELECTION_FREEZE)
            & (frame[f'target_available_date_assumed_h{h}'] < SELECTION_FREEZE))
        frame[f'origin_period_test_eligible_h{h}'] = (
            complete & frame.origin_period_group.eq('disruption_test'))
        frame[f'forecast_issue_test_eligible_assumed_h{h}'] = (
            complete & frame.forecast_issue_period_group_assumed.eq('disruption_test'))
        frame[f'strict_point_in_time_verified_h{h}'] = (
            complete & frame.historical_release_dates_verified & frame.historical_data_vintages_verified)
    return frame


def verify(frame, freight, brent):
    check_months(frame, 'Output')
    if not frame.feature_complete.all():
        raise ValueError('Expected complete features after using 2004 warm-up')
    if np.isinf(frame.select_dtypes(include=[np.number]).to_numpy()).any():
        raise ValueError('Infinite output values')
    raw_freight = freight.set_index('date').freight_index
    raw_brent = brent.set_index('date').brent_usd_per_barrel
    assert np.allclose(frame.freight_index, raw_freight.reindex(frame.date))
    assert np.allclose(frame.brent_usd_per_barrel, raw_brent.reindex(frame.date))
    for lag in (1, 2, 3, 6, 12):
        expected = np.log(raw_freight.reindex(frame.date).to_numpy() / raw_freight.reindex(frame.date-pd.DateOffset(months=lag)).to_numpy())
        assert np.allclose(frame[f'freight_log_change_{lag}m'], expected)
    assert np.allclose(frame.brent_log_change_1m, np.log(raw_brent.reindex(frame.date).to_numpy() / raw_brent.reindex(frame.date-pd.DateOffset(months=1)).to_numpy()))
    for window in (6, 12):
        independent = []
        for date in frame.date:
            months = pd.date_range(date-pd.DateOffset(months=window-1), date, freq='MS')
            changes = np.log(raw_freight.reindex(months).to_numpy() / raw_freight.reindex(months-pd.DateOffset(months=1)).to_numpy())
            independent.append(np.std(changes, ddof=0))
        assert np.allclose(frame[f'freight_volatility_{window}m'], independent)
    for h in HORIZONS:
        observed = frame[f'target_observed_h{h}']
        expected = raw_freight.reindex(frame.loc[observed, f'target_month_h{h}']).to_numpy()
        assert np.allclose(frame.loc[observed, f'target_freight_index_h{h}'], expected)
        assert np.allclose(frame.loc[observed, f'target_freight_log_change_h{h}'], np.log(expected/frame.loc[observed, 'freight_index']))
        assert frame.loc[~observed, f'target_freight_log_change_h{h}'].isna().all()
        assert (frame.loc[frame[f'validation_selection_eligible_h{h}'], f'target_available_date_assumed_h{h}'] < SELECTION_FREEZE).all()
        assert (frame.loc[frame[f'initial_training_fit_eligible_h{h}'], f'target_available_date_assumed_h{h}'] <= frame.initial_validation_asof_date_assumed.iloc[0]).all()
    assert not any(c.startswith('target_') for c in FEATURES)
    return dict(raw_values_reconciled=True,month_gaps_checked=True,lag_features_recomputed=True,
                volatility_recomputed=True,future_targets_recomputed=True,selection_cutoff_checked=True,
                future_missing_targets_retained=True,source_csv_not_modified=True)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if OUTPUT.exists():
        raise FileExistsError(f'Existing training data preserved: {OUTPUT}')
    freight, brent = read_freight(), read_brent()
    frame = prepare(freight, brent)
    checks = verify(frame, freight, brent)
    sources = []
    for label,path,url in [
        ('BLS/FRED freight index',FREIGHT_PATH,'https://fred.stlouisfed.org/series/PCU483111483111'),
        ('World Bank Brent',BRENT_PATH,'https://www.worldbank.org/en/research/commodity-markets'),
    ]:
        sources.append(dict(source=label,file=str(path),url=url,sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            file_modified_utc=datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(),
            download_date_verified=False,
            note='File modification time is provenance metadata, not a verified download or historical publication date.'))
    summary = dict(
        prepared_on_client_date=PREPARED_DATE,output_file=str(OUTPUT),rows=len(frame),columns=len(frame.columns),
        start=frame.date.min().strftime('%Y-%m'),end=frame.date.max().strftime('%Y-%m'),
        feature_columns=FEATURES,quantile_outputs_required=[.05,.10,.25,.50,.75,.90,.95],
        horizons=list(HORIZONS),h12_reporting_only=True,
        target_definition='ln(freight_index[t+h] / freight_index[t]); t is reference month of latest included observation',
        split_basis='origin reference month; forecast issue dates recorded separately',
        requested_split=dict(initial_training='2005-2014',validation='2015-2019',disruption_test='2020-2021',post_disruption='2022 onward'),
        origin_period_counts={k:int(v) for k,v in frame.origin_period_group.value_counts().items()},
        horizon_counts={f'h{h}':dict(targets_observed=int(frame[f'target_observed_h{h}'].sum()),
            initial_fit_rows=int(frame[f'initial_training_fit_eligible_h{h}'].sum()),
            validation_rows_with_known_pretest_targets=int(frame[f'validation_selection_eligible_h{h}'].sum()),
            disruption_origin_rows=int(frame[f'origin_period_test_eligible_h{h}'].sum()),
            strict_point_in_time_verified_rows=int(frame[f'strict_point_in_time_verified_h{h}'].sum())) for h in HORIZONS},
        warmup_policy='2004 used only to construct lag/volatility features. Export and labelled history start January 2005.',
        volatility_definition='Population standard deviation (ddof=0) of latest 6/12 monthly freight log returns, including current month; not annualized.',
        availability_policy='ASSUMED: observations usable at end of following calendar month; not a verified release calendar.',
        prototype_status='Prepared research snapshot. Verified historical release dates and revision vintages remain incomplete.',
        split_caution='User selected 2015-2019 main validation; specification mentions 2012-2019. Optional early-backtest flag retains 2012-2014.',
        selection_freeze_date=str(SELECTION_FREEZE.date()),
        remaining_inputs=['Historical release-date records / vintages','FBX 2020-2021 comparison data'],
        sources=sources,verification=checks,
    )
    output = frame.copy()
    for column in output:
        if pd.api.types.is_datetime64_any_dtype(output[column]):
            output[column] = output[column].dt.strftime('%Y-%m-%d')
    output.to_csv(OUTPUT,index=False,float_format='%.12g',na_rep='')
    saved = pd.read_csv(OUTPUT)
    assert list(saved.columns) == list(frame.columns) and len(saved) == len(frame)
    assert saved[FEATURES].notna().all().all()
    for source in sources:
        assert hashlib.sha256(Path(source['file']).read_bytes()).hexdigest() == source['sha256']
    (OUT_DIR/'model_c_preparation_manifest.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
    print(json.dumps({k:summary[k] for k in ['output_file','rows','columns','start','end','origin_period_counts','horizon_counts']},indent=2))


if __name__ == '__main__':
    main()
